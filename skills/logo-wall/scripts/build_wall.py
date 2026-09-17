#!/usr/bin/env python3
"""
Build a logo-wall slide (editable PPTX) from a wall.json config.

    python3 build_wall.py wall.json out.pptx [--check]

--check downloads and pixel-verifies every brand's candidates and prints the
confirm report WITHOUT building. Run it first, show the report to the user,
build only after they confirm — flags are decisions for a human, not the tool.

wall.json:
{
  "title": "Trusted by teams everywhere",
  "background": "dark",                  // "dark" | "light" | 6-digit hex
  "mono": true,                          // require monochrome marks
  "cols": null,                          // null = derived from brand count
  "brands": [
    { "name": "Slack",
      "candidates": ["<primary src URL>", "<type:other fallback URL>", ...] }
  ]
}

Candidate URLs come from get_brand responses — the theme-matched primary
first, then any same-theme variants under type "other". The script downloads
candidates in order and keeps the FIRST one that passes the pixel check
(when mono is requested). The `theme` label names the asset's own colour but
does NOT guarantee monochrome: official light-theme variants often keep a
coloured symbol, and the true all-white version, when it exists, may be
filed under type "other" with no semantic label. Only pixels settle it.

This script never recolours a logo. A brand with no passing variant is
FLAGGED, and the user chooses: accept the least-coloured candidate, switch
that brand to its symbol, or drop it. Force-tinting a brand's mark is a
brand-guideline violation, not a fallback.
"""
import json
import math
import os
import re
import subprocess
import sys
import warnings

warnings.filterwarnings("ignore")

from PIL import Image

DARK = "101012"
LIGHT = "FAFAF8"
SAT_THRESH = 0.15     # per-pixel saturation above this counts as coloured
FRAC_THRESH = 0.02    # >2% coloured visible pixels = not monochrome
CLAMP_HI = 1.25       # natural-height band around the median
CLAMP_LO = 0.72

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN, TOP, BOTTOM = 0.9, 1.4, 0.7


def fetch(url, dest):
    os.makedirs(os.path.dirname(dest) or ".", exist_ok=True)
    subprocess.run(["curl", "-sS", "--fail", "-o", dest, url],
                   check=True, capture_output=True)
    return dest


def coloured_fraction(path):
    """Fraction of visible pixels with meaningful saturation."""
    im = Image.open(path).convert("RGBA")
    im.thumbnail((200, 200))
    visible = coloured = 0
    for r, g, b, a in im.getdata():
        if a <= 40:
            continue
        visible += 1
        mx, mn = max(r, g, b), min(r, g, b)
        if mx > 30 and (mx - mn) / mx > SAT_THRESH:
            coloured += 1
    return (coloured / visible) if visible else 0.0


def ink_ratio(path):
    """Aspect ratio of the tight ink bounding box (not the canvas)."""
    im = Image.open(path).convert("RGBA")
    bb = im.getbbox()
    if not bb:
        return 1.0
    return (bb[2] - bb[0]) / max(1, bb[3] - bb[1])


def resolve_brand(brand, workdir, want_mono):
    """
    Download candidates in order; return (path, flag_or_None).
    Keeps the first mono pass when mono is requested, else the first candidate.
    """
    slug = re.sub(r"[^a-z0-9]+", "-", brand["name"].lower()).strip("-")
    best, best_frac = None, 1.0
    for i, url in enumerate(brand["candidates"]):
        path = os.path.join(workdir, f"{slug}-{i}.png")
        try:
            fetch(url, path)
            Image.open(path).verify()
        except Exception as e:
            continue
        frac = coloured_fraction(path)
        if not want_mono:
            return path, None
        if frac < FRAC_THRESH:
            return path, None
        if frac < best_frac:
            best, best_frac = path, frac
    if best is None:
        return None, f"{brand['name']}: no candidate downloaded successfully"
    return best, (f"{brand['name']}: no fully monochrome variant — best candidate "
                  f"is {best_frac*100:.1f}% coloured. Accept it, switch this brand "
                  f"to its symbol, or drop it. (Never force-tint.)")


def grid_shape(n, cols):
    if cols:
        return cols, math.ceil(n / cols)
    cols = 3 if n <= 6 else 4 if n <= 12 else 5 if n <= 20 else 6
    return cols, math.ceil(n / cols)


def build(cfg, out, check_only=False):
    bg = cfg.get("background", "dark")
    bg_hex = DARK if bg == "dark" else LIGHT if bg == "light" else bg.lstrip("#").upper()
    want_mono = bool(cfg.get("mono", True))
    workdir = "wall_assets"

    resolved, flags = [], []
    for brand in cfg["brands"]:
        path, flag = resolve_brand(brand, workdir, want_mono)
        if flag:
            flags.append(flag)
        if path:
            resolved.append({"name": brand["name"], "path": path,
                             "ratio": ink_ratio(path)})

    cols, rows = grid_shape(len(resolved), cfg.get("cols"))
    print(f"brands resolved: {len(resolved)}/{len(cfg['brands'])}  "
          f"grid: {cols} x {rows}  background: {bg_hex}  mono: {want_mono}")
    for f in flags:
        print(f"  ! {f}")
    if check_only:
        print("\n--check: nothing built. Confirm the report, then run without --check.")
        return 1 if flags else 0

    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(bg_hex)

    if cfg.get("title"):
        from pptx.enum.text import PP_ALIGN
        tb = slide.shapes.add_textbox(Inches(MARGIN), Inches(0.55),
                                      Inches(SLIDE_W - 2 * MARGIN), Inches(0.6))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = cfg["title"]
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.name = "Arial"
        dark_bg = int(bg_hex[:2], 16) * 0.3 + int(bg_hex[2:4], 16) * 0.6 < 100
        run.font.color.rgb = RGBColor.from_string(LIGHT if dark_bg else "1A1A18")

    grid_w, grid_h = SLIDE_W - 2 * MARGIN, SLIDE_H - TOP - BOTTOM
    cell_w, cell_h = grid_w / cols, grid_h / rows
    box_w, box_h = cell_w * 0.66, cell_h * 0.42

    # pass 1: natural contained size per logo
    for r in resolved:
        w = box_w
        h = w / r["ratio"]
        if h > box_h:
            h, w = box_h, box_h * r["ratio"]
        r["w"], r["h"] = w, h

    # pass 2: clamp heights to a band around the median — vary within bounds.
    # Naive sizing looks organic when the set cooperates; the clamp is what
    # keeps one square symbol or one skyscraper glyph from wrecking the wall.
    heights = sorted(r["h"] for r in resolved)
    med = heights[len(heights) // 2]
    hi, lo = med * CLAMP_HI, med * CLAMP_LO
    for r in resolved:
        h, w = r["h"], r["w"]
        if h > hi:
            h, w = hi, hi * r["ratio"]
        elif h < lo:
            h, w = lo, lo * r["ratio"]
            if w > box_w:
                w, h = box_w, box_w / r["ratio"]
        r["w"], r["h"] = w, h

    # place; centre a partial last row
    for i, r in enumerate(resolved):
        row, col = divmod(i, cols)
        in_row = cols if row < rows - 1 else len(resolved) - (rows - 1) * cols
        row_offset = (cols - in_row) * cell_w / 2 if row == rows - 1 else 0
        cx = MARGIN + row_offset + col * cell_w
        cy = TOP + row * cell_h
        slide.shapes.add_picture(r["path"],
                                 Inches(cx + (cell_w - r["w"]) / 2),
                                 Inches(cy + (cell_h - r["h"]) / 2),
                                 Inches(r["w"]), Inches(r["h"]))

    prs.save(out)
    print(f"\nwritten: {out}")
    if flags:
        print("shipped WITH flags above — make sure the user accepted them.")
    return 0


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) < 1:
        sys.exit("usage: build_wall.py wall.json [out.pptx] [--check]")
    cfg = json.load(open(args[0]))
    out = args[1] if len(args) > 1 else "wall.pptx"
    sys.exit(build(cfg, out, check_only="--check" in sys.argv))
